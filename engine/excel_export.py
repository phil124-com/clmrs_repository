"""
Builds the client-facing Excel workbook, matching the visual style of
Philemon's standing household-profiling reports: navy Arial titles, italic
gray methodology/subtitle text under each title, white-on-navy table headers,
freeze panes, autofilter tables, and one row per unique farmer/child combining
raw and analytical fields — applied consistently across Household, Inspection,
and Follow Up.
"""
from __future__ import annotations
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

NAVY = "1F4E78"
HEADER_FILL = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
KPI_FILL = PatternFill(start_color="E7EEF6", end_color="E7EEF6", fill_type="solid")
TITLE_FONT = Font(name="Arial", size=14, bold=True, color=NAVY)
SUMMARY_TITLE_FONT = Font(name="Arial", size=16, bold=True, color=NAVY)
SUBTITLE_FONT = Font(name="Arial", size=10, italic=True, color="555555")
NOTE_FONT = Font(name="Arial", size=9, italic=True, color="777777")
HEADER_FONT = Font(name="Arial", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=10)
KPI_LABEL_FONT = Font(name="Arial", size=11, bold=True)
KPI_VALUE_FONT = Font(name="Arial", size=16, bold=True, color=NAVY)


def _fmt_date(d):
    if pd.isna(d):
        return None
    return d.strftime("%Y-%m-%d") if hasattr(d, "strftime") else d


def _write_table(ws, df: pd.DataFrame, start_row: int, table_name: str, freeze: bool = True) -> int:
    if df is None or df.empty:
        ws.cell(row=start_row, column=1, value="(no records)").font = BODY_FONT
        return start_row + 2

    for j, col in enumerate(df.columns, start=1):
        c = ws.cell(row=start_row, column=j, value=str(col))
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for i, (_, row) in enumerate(df.iterrows(), start=start_row + 1):
        for j, col in enumerate(df.columns, start=1):
            val = row[col]
            if pd.isna(val):
                val = None
            elif isinstance(val, pd.Timestamp):
                val = val.to_pydatetime()
            cell = ws.cell(row=i, column=j, value=val)
            cell.font = BODY_FONT
            if hasattr(val, "strftime"):
                cell.number_format = "yyyy-mm-dd"

    last_row = start_row + len(df)
    last_col = len(df.columns)
    for j, col in enumerate(df.columns, start=1):
        sample = df[col].fillna("").astype(str).values[:300]
        max_len = max([len(str(col))] + [len(str(s)) for s in sample])
        ws.column_dimensions[get_column_letter(j)].width = min(max(max_len + 2, 10), 50)

    ws.freeze_panes = ws.cell(row=start_row + 1, column=1) if freeze else None

    safe_name = "".join(ch if ch.isalnum() or ch == "_" else "_" for ch in table_name)[:30]
    ref = f"{get_column_letter(1)}{start_row}:{get_column_letter(last_col)}{last_row}"
    try:
        tab = Table(displayName=safe_name, ref=ref)
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tab)
    except Exception:
        pass

    return last_row + 2


def _title_block(ws, title: str, subtitle: str = None, note: str = None, width: int = 10) -> int:
    r = 1
    ws.cell(row=r, column=1, value=title).font = TITLE_FONT
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=width)
    r += 1
    if subtitle:
        ws.cell(row=r, column=1, value=subtitle).font = SUBTITLE_FONT
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=width)
        r += 1
    if note:
        ws.cell(row=r, column=1, value=note).font = NOTE_FONT
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=width)
        r += 1
    return r + 1  # blank spacer row


def _kpi_grid(ws, start_row: int, kpis: list) -> int:
    """kpis: list of (label, value). Rendered as a row of labelled boxes,
    4 per row, matching the Summary sheet's KPI-tile look."""
    per_row = 4
    r = start_row
    for i in range(0, len(kpis), per_row):
        chunk = kpis[i:i + per_row]
        for j, (label, value) in enumerate(chunk):
            col = j * 2 + 1
            lc = ws.cell(row=r, column=col, value=label)
            lc.font = KPI_LABEL_FONT
            lc.fill = KPI_FILL
            ws.merge_cells(start_row=r, start_column=col, end_row=r, end_column=col + 1)
            vc = ws.cell(row=r + 1, column=col, value=value)
            vc.font = KPI_VALUE_FONT
            vc.fill = KPI_FILL
            vc.alignment = Alignment(horizontal="center")
            ws.merge_cells(start_row=r + 1, start_column=col, end_row=r + 1, end_column=col + 1)
        r += 3
    for col in range(1, per_row * 2 + 1):
        ws.column_dimensions[get_column_letter(col)].width = 16
    return r + 1


def build_report(sheets: dict, exec_summary: dict, warnings: list, programme_name: str,
                  meta: dict, output_path: str):
    wb = Workbook()
    wb.remove(wb.active)

    # ---------------------------------------------------------------- Summary
    ws = wb.create_sheet("Summary")
    groups = ", ".join(meta.get("farmer_groups", []) or [])
    date_range = ""
    if meta.get("date_min") is not None and meta.get("date_max") is not None:
        date_range = f"Data covers {_fmt_date(meta['date_min'])} to {_fmt_date(meta['date_max'])}"
    subtitle = "  |  ".join(x for x in [
        f"Farmer groups: {groups}" if groups else "",
        date_range,
        f"Validated against Annex farmer list ({meta.get('farmer_list_file', '')}, {meta.get('farmer_list_count', 0):,} farmers)" if meta.get("farmer_list_file") else "",
    ] if x)
    r = _title_block(ws, f"{programme_name} Programme \u2014 CLMRS Analysis Summary", subtitle, width=8)

    kpis = [(k, v) for k, v in exec_summary.items()]
    r = _kpi_grid(ws, r, kpis)

    if not sheets.get("hh_match_by_group", pd.DataFrame()).empty:
        ws.cell(row=r, column=1, value="Farmer List (Annex) Match Rate by Farmer Group \u2014 Household").font = Font(bold=True, size=11)
        r += 1
        ws.cell(row=r, column=1,
                value="A single overall percentage can hide big differences between groups \u2014 shown here so a low match rate isn't missed at just one group.").font = NOTE_FONT
        r += 1
        r = _write_table(ws, sheets["hh_match_by_group"], r, "MatchByGroupHH", freeze=False)

    if not sheets.get("insp_match_by_group", pd.DataFrame()).empty:
        ws.cell(row=r, column=1, value="Farmer List (Annex) Match Rate by Farmer Group \u2014 Inspection").font = Font(bold=True, size=11)
        r += 1
        r = _write_table(ws, sheets["insp_match_by_group"], r, "MatchByGroupInsp", freeze=False)

    if warnings:
        ws.cell(row=r, column=1, value="Data Load Warnings").font = Font(bold=True, size=11)
        r += 1
        for w in warnings:
            ws.cell(row=r, column=1, value=f"\u2022 {w}").font = BODY_FONT
            r += 1

    # ---------------------------------------------------------------- Household
    ws = wb.create_sheet("HH - Farmers by Quarter")
    r = _title_block(ws, "Household \u2014 Unique Farmers Profiled, by Quarter (On Annex List)",
                      "One row per unique farmer (Farmer ID) matching the Annex farmer list. \"Children Profiled\" is the total unique children ever recorded under that farmer. Farmer placed in the quarter of their FIRST profiling visit.",
                      "Farmers NOT on the Annex list are on the \"HH - Not on List\" tab.", width=10)
    _write_table(ws, sheets["hh_farmer_on_list"], r, "HHFarmersOnList")

    ws = wb.create_sheet("HH - Children Detail")
    r = _title_block(ws, "Household \u2014 Unique Children Profiled (On Annex List)",
                      "One row per unique child (Child ID) whose farmer matches the Annex farmer list, combining raw survey fields with CLMRS status, school attendance, and education level.",
                      "Children whose farmer is not on the Annex list are on the \"HH - Not on List\" tab.", width=11)
    _write_table(ws, sheets["hh_child_on_list"], r, "HHChildrenOnList")

    ws = wb.create_sheet("HH - Not on List")
    r = _title_block(ws, "Household \u2014 Farmers & Children NOT on the Annex List",
                      "Profiled in the field but the Farmer ID does not match any ID in the Annex farmer list. Excluded from all main Household analysis tabs and summaries; kept here for follow-up/data cleaning.",
                      width=10)
    ws.cell(row=r, column=1, value="Farmers Not on List").font = Font(bold=True, size=11)
    r += 1
    r = _write_table(ws, sheets["hh_farmer_not_on_list"], r, "HHFarmersOffList", freeze=False)
    ws.cell(row=r, column=1, value="Children Not on List").font = Font(bold=True, size=11)
    r += 1
    r = _write_table(ws, sheets["hh_child_not_on_list"], r, "HHChildrenOffList", freeze=False)

    ws = wb.create_sheet("HH - Quarterly & Status")
    r = _title_block(ws, "Household \u2014 Quarterly, Status, Gender & School Summary (Annex-Matched Only)",
                      "Pivoted from the Household detail tabs above. Counts cover only farmers/children present on the Annex list.", width=6)
    hh_q = sheets["hh_quarterly"]
    ws.cell(row=r, column=1, value="Unique Farmers & Children Profiled, by Quarter").font = Font(bold=True, size=11)
    r += 1
    r = _write_table(ws, hh_q["summary"], r, "HHQuarterly", freeze=False)
    ws.cell(row=r, column=1,
            value=f"Farmers with a repeat visit in a later quarter (still counted once, in their first quarter): {hh_q['repeat_visit_farmers']}").font = NOTE_FONT
    r += 2
    ws.cell(row=r, column=1, value="CLMRS Case Status, by Quarter").font = Font(bold=True, size=11)
    r += 1
    r = _write_table(ws, sheets["hh_status"]["pivot"], r, "HHStatus", freeze=False)
    ws.cell(row=r, column=1,
            value=f"Farmer households with no child under 18 (Annex-matched, all quarters): {sheets['hh_status']['no_child_farmers']}").font = NOTE_FONT
    r += 2
    ws.cell(row=r, column=1, value="Children by Gender, by Quarter").font = Font(bold=True, size=11)
    r += 1
    r = _write_table(ws, sheets["hh_child_gender"], r, "HHChildGender", freeze=False)
    ws.cell(row=r, column=1, value="Farmers by Gender (from Farmer List / Annex)").font = Font(bold=True, size=11)
    r += 1
    r = _write_table(ws, sheets["hh_farmer_gender"], r, "HHFarmerGender", freeze=False)
    ws.cell(row=r, column=1, value="Records by Enumerator Gender").font = Font(bold=True, size=11)
    r += 1
    r = _write_table(ws, sheets["hh_enum_gender"], r, "HHEnumGender", freeze=False)
    r += 1
    ws.cell(row=r, column=1, value="School Attendance, by Quarter").font = Font(bold=True, size=11)
    r += 1
    r = _write_table(ws, sheets["hh_school"]["attendance"], r, "HHSchool", freeze=False)
    ws.cell(row=r, column=1, value="Current Education Level (all quarters)").font = Font(bold=True, size=11)
    r += 1
    r = _write_table(ws, sheets["hh_school"]["education"], r, "HHEducation", freeze=False)

    # ---------------------------------------------------------------- Inspection
    ws = wb.create_sheet("Insp - Farmers by Quarter")
    r = _title_block(ws, "Inspection \u2014 Unique Farmers Inspected, by Quarter (On Annex List)",
                      "One row per unique farmer (Farmer ID) matching the Annex farmer list, inspected during the period covered. Same unique-ID methodology as Household.",
                      "Farmers NOT on the Annex list are on the \"Insp - Not on List\" tab.", width=10)
    r = _write_table(ws, sheets["insp_farmer_on_list"], r, "InspFarmersOnList")
    ws.cell(row=r, column=1, value="Farmers Inspected by Gender (from Farmer List / Annex)").font = Font(bold=True, size=11)
    r += 1
    r = _write_table(ws, sheets["insp_farmer_gender"], r, "InspFarmerGender", freeze=False)
    ws.cell(row=r, column=1, value="Records by Enumerator Gender").font = Font(bold=True, size=11)
    r += 1
    r = _write_table(ws, sheets["insp_enum_gender"], r, "InspEnumGender", freeze=False)

    ws = wb.create_sheet("Insp - Children Detail")
    r = _title_block(ws, "Inspection \u2014 Unique Children Inspected (On Annex List)",
                      "One row per unique child (Child ID) whose farmer matches the Annex farmer list, combining raw survey fields with CLMRS status, gender, and school attendance.",
                      "Children whose farmer is not on the Annex list are on the \"Insp - Not on List\" tab.", width=11)
    r = _write_table(ws, sheets["insp_child_on_list"], r, "InspChildrenOnList")
    ws.cell(row=r, column=1, value="Children Inspected by Gender, by Quarter").font = Font(bold=True, size=11)
    r += 1
    r = _write_table(ws, sheets["insp_child_gender"], r, "InspChildGender", freeze=False)

    ws = wb.create_sheet("Insp - Not on List")
    r = _title_block(ws, "Inspection \u2014 Farmers & Children NOT on the Annex List", width=10)
    ws.cell(row=r, column=1, value="Farmers Not on List").font = Font(bold=True, size=11)
    r += 1
    r = _write_table(ws, sheets["insp_farmer_not_on_list"], r, "InspFarmersOff", freeze=False)
    ws.cell(row=r, column=1, value="Children Not on List").font = Font(bold=True, size=11)
    r += 1
    r = _write_table(ws, sheets["insp_child_not_on_list"], r, "InspChildrenOff", freeze=False)

    ws = wb.create_sheet("HH vs Inspection")
    r = _title_block(ws, "Household vs Inspection \u2014 Monitoring Coverage",
                      "Classifies every unique farmer/child as profiled-and-inspected, profiled-but-not-inspected, or inspected-but-not-profiled, to surface monitoring coverage gaps.", width=4)
    ws.cell(row=r, column=1, value="Farmers").font = Font(bold=True, size=11)
    r += 1
    r = _write_table(ws, sheets["hh_vs_insp"]["farmers"], r, "HHvsInspFarmers", freeze=False)
    ws.cell(row=r, column=1, value="Children").font = Font(bold=True, size=11)
    r += 1
    r = _write_table(ws, sheets["hh_vs_insp"]["children"], r, "HHvsInspChildren", freeze=False)

    # ---------------------------------------------------------------- Follow Up
    ws = wb.create_sheet("Follow Up Cases")
    r = _title_block(ws, "Follow Up \u2014 Unique Child Labour Cases",
                      "One row per unique Child ID in the Follow Up dataset. Status and Case Identified Date reflect the LATEST follow-up record; Quarter reflects the FIRST follow-up visit for that child (a child can be followed up more than once before the case closes).",
                      "Case Identified Date = Follow Up Date minus 3 calendar months (Pass/Fail) or minus 6 calendar months (Fully Remediated).", width=10)
    ws.cell(row=r, column=1, value="Case Status Summary").font = Font(bold=True, size=11)
    r += 1
    r = _write_table(ws, sheets["fu_summary"], r, "FUSummary", freeze=False)
    ws.cell(row=r, column=1, value="Case Detail").font = Font(bold=True, size=11)
    r += 1
    r = _write_table(ws, sheets["fu_cases"], r, "FUCases", freeze=False)

    ws = wb.create_sheet("FU - Multi-Visit Timeline")
    r = _title_block(ws, "Follow Up \u2014 Multi-Visit Timeline",
                      "For every child in the Follow Up dataset: the number of follow-up visits, the gap in days between consecutive visits, and \u2014 for children who reach Fully Remediated \u2014 the time from their FIRST follow-up to the remediation date.",
                      "The event log below lists every individual visit; the summary above condenses it to one row per child.", width=9)
    ws.cell(row=r, column=1, value="Per-Child Summary").font = Font(bold=True, size=11)
    r += 1
    r = _write_table(ws, sheets["fu_timeline_summary"], r, "FUTimelineSummary", freeze=False)
    ws.cell(row=r, column=1, value="Visit-Level Event Log").font = Font(bold=True, size=11)
    r += 1
    r = _write_table(ws, sheets["fu_timeline_events"], r, "FUTimelineEvents", freeze=False)

    ws = wb.create_sheet("Follow Up Alerts")
    r = _title_block(ws, "Follow Up Alerts \u2014 Next Visit Due",
                      "Children whose latest Follow Up status is neither \"Follow up pass\" nor \"Fully remediated\" \u2014 i.e. the case is still open and another follow-up visit is needed.",
                      "Next Follow-Up Due = last follow-up date + 3 calendar months. \"OVERDUE\" means that date has already passed as of when this report was generated; \"Due Soon\" means it hasn't yet.", width=8)
    _write_table(ws, sheets["fu_alerts"], r, "FUAlerts")

    ws = wb.create_sheet("No Children (HH + Insp)")
    r = _title_block(ws, "Farmers with No Child Under 18 \u2014 Household + Inspection Combined",
                      "Every farmer flagged as having no child under 18, from BOTH the Household and Inspection datasets, checked against the Annex farmer list. A farmer visited under both datasets appears once per source below.", width=7)
    ws.cell(row=r, column=1, value="Unique Farmer Summary").font = Font(bold=True, size=11)
    r += 1
    r = _write_table(ws, sheets["no_children_combined_summary"], r, "NoChildrenSummary", freeze=False)
    ws.cell(row=r, column=1, value="Detail").font = Font(bold=True, size=11)
    r += 1
    r = _write_table(ws, sheets["no_children_combined_detail"], r, "NoChildrenDetail", freeze=False)

    # ---------------------------------------------------------------- Enumerator & Data Quality
    ws = wb.create_sheet("Enumerator Analysis")
    r = _title_block(ws, "Enumerator Analysis",
                      "Unique farmers/children per enumerator, across Household, Inspection, and Follow Up.", width=7)
    _write_table(ws, sheets["enumerator"], r, "Enumerator")

    ws = wb.create_sheet("Data Quality")
    r = _title_block(ws, "Data Quality Exceptions",
                      "Genuine data issues only \u2014 fields that legitimately don't apply (e.g. a farmer visit with no child) are excluded from these counts.", width=4)
    _write_table(ws, sheets["data_quality"], r, "DataQuality")

    wb.save(output_path)
    return output_path
